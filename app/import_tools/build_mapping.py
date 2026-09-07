# -*- coding: utf-8 -*-
"""掃「2025桂圓銷售_整理版.xlsx」,產出「匯入對照表.xlsx」。
表裡先填好『建議』欄,留空『確認』欄給郡碩農創修正。

用法:
    python build_mapping.py                     # 讀 ../../2025桂圓銷售_整理版.xlsx
    python build_mapping.py 路徑\整理版.xlsx     # 指定來源
輸出:整理版同資料夾下的 匯入對照表.xlsx
"""
import sys, os, re, datetime as dt
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "..", "2025桂圓銷售_整理版.xlsx")
SRC = os.path.abspath(SRC)
OUT = os.path.join(os.path.dirname(SRC), "匯入對照表.xlsx")

# 整理版「銷售明細」欄位索引
C_SEASON, C_YM, C_IDX, C_DATE, C_CUST, C_CSRC, C_ITEM, C_SPEC = 0, 1, 2, 3, 4, 5, 6, 7
C_QTY, C_PRICE, C_TOTAL = 10, 11, 12
C_KIND, C_PAY, C_BANK, C_PLAN = 14, 15, 16, 17
C_LOGI, C_CARRIER, C_SELFLOC = 21, 22, 23
C_CHECK = 26

# ---- 讀來源 --------------------------------------------------------
if not os.path.exists(SRC):
    sys.exit(f"找不到來源檔:{SRC}")
wb = openpyxl.load_workbook(SRC, data_only=True)
if "銷售明細" not in wb.sheetnames:
    sys.exit("來源檔沒有「銷售明細」分頁")
rows = [r for r in wb["銷售明細"].iter_rows(min_row=2, values_only=True) if r[C_ITEM]]
print(f"讀入 {len(rows)} 列(有品項)")


# ---- 猜測邏輯 ----------------------------------------------------
def guess_sku(item, spec):
    s = (spec or "").strip()
    meat = "肉" in (item or "")
    if "1000" in s:
        return "LG-MEAT-1000", "龍眼肉 1000g", "單品"
    if "600" in s:
        return ("LG-MEAT-600", "龍眼肉 600g 罐", "單品") if meat else ("GY-GIFT", "桂圓乾 禮盒", "禮盒")
    if "500" in s:
        return "GY-DRY-500", "桂圓乾 500g", "單品"
    if "300" in s:
        return "GY-DRY-300", "桂圓乾 300g", "單品"
    if "200" in s:
        return "LG-MEAT-200", "龍眼肉 200g", "單品"
    if any(k in s for k in ("裸", "散", "真空")) or s in ("#3", ""):
        return ("LG-MEAT-BULK", "龍眼肉 裸裝", "裸裝") if meat else ("GY-DRY-BULK", "桂圓乾 裸裝", "裸裝")
    if meat:
        return "LG-MEAT-600", "龍眼肉", "單品"
    return "GY-DRY-300", "桂圓乾", "單品"


def split_name(cust):
    n = (cust or "").replace("\t", "").strip()
    m = re.match(r"^(.*?)[（(](.+?)[）)]\s*$", n)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    if "/" in n:
        a, b = n.split("/", 1)
        return a.strip(), b.strip().rstrip("/")
    return n, ""


def guess_segment(cust, csrc, kinds):
    blob = f"{cust} {csrc}"
    if "公關" in kinds and not any(k != "公關" for k in kinds):
        return "公關對象"
    if any(k in blob for k in ("米糧行", "有機商店", "聯盟", "商行", "企業", "公司")):
        return "批發"
    if any(k in blob for k in ("市集", "育成", "團媽", "團購", "老師")):
        return "團購主"
    if any(k in blob for k in ("今周刊", "微笑台灣", "研究院", "縣府", "縣政府", "科大", "雲科", "職訓", "SBIR", "國發會", "台經院")):
        return "機構"
    return "零售"


def guess_channel(csrc, logi, plan):
    blob = f"{csrc} {logi} {plan}"
    if "官網" in blob:
        return "官網"
    if "LINE" in blob or "line" in blob:
        return "LINE社群"
    if any(k in blob for k in ("台灣亮起來", "今周刊", "微笑台灣", "點石成金")):
        return "媒體導流"
    if any(k in blob for k in ("市集", "育成", "成果展", "博覽會", "健行節")):
        return "市集展售"
    if any(k in blob for k in ("米糧行", "有機商店", "一品園", "批發")):
        return "批發"
    if logi in ("自送", "自取") or "自送" in (logi or ""):
        return "自售"
    return "(待定)"


# ---- 彙整 ------------------------------------------------------
prod = Counter()
prod_price = defaultdict(list)
cust_stat = defaultdict(lambda: {"n": 0, "amt": 0.0, "csrc": Counter(), "kinds": Counter()})
csrc_logi = Counter()
season_cnt = Counter()

for r in rows:
    item, spec = (r[C_ITEM] or "").strip(), (r[C_SPEC] or "").strip() if r[C_SPEC] else ""
    prod[(item, spec)] += 1
    if r[C_KIND] == "銷售" and r[C_PRICE]:
        prod_price[(item, spec)].append(r[C_PRICE])
    cust = (r[C_CUST] or "").replace("\t", "").strip()
    st = cust_stat[cust]
    st["n"] += 1
    st["amt"] += (r[C_TOTAL] or 0) if r[C_KIND] == "銷售" else 0
    if r[C_CSRC]:
        st["csrc"][r[C_CSRC]] += 1
    st["kinds"][r[C_KIND] or "銷售"] += 1
    csrc_logi[((r[C_CSRC] or "").strip(), (r[C_LOGI] or "").strip(), (r[C_PLAN] or "").strip())] += 1
    season_cnt[r[C_SEASON]] += 1

# ---- 寫出對照表 --------------------------------------------------
wb2 = openpyxl.Workbook()
wb2.remove(wb2.active)
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4F6228")
NEED = PatternFill("solid", fgColor="FFF2CC")   # 要填的欄底色


def sheet(name, headers, data, widths, need_cols):
    ws = wb2.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = HDR; c.fill = FILL; c.alignment = Alignment(horizontal="center")
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for ci in need_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=ci).fill = NEED
    ws.freeze_panes = "A2"
    return ws


# 商品對照
data = []
for (item, spec), n in sorted(prod.items(), key=lambda x: -x[1]):
    sku, pname, ptype = guess_sku(item, spec)
    pr = prod_price[(item, spec)]
    common = Counter(pr).most_common(1)[0][0] if pr else ""
    data.append([item, spec or "(空)", n, sku, pname, ptype, common, "", ""])
sheet("商品對照",
      ["整理版品項", "整理版規格", "出現次數", "建議SKU", "建議品名", "建議型態",
       "最常見成交單價", "★確認SKU", "備註"],
      data, [12, 14, 9, 16, 20, 10, 14, 16, 24], [8, 9])

# 客戶對照
data = []
for cust, st in sorted(cust_stat.items(), key=lambda x: -x[1]["amt"]):
    nm, src = split_name(cust)
    seg = guess_segment(cust, " ".join(st["csrc"]), st["kinds"])
    csrc_top = st["csrc"].most_common(1)[0][0] if st["csrc"] else ""
    data.append([cust, csrc_top, st["n"], round(st["amt"]), nm, seg, "", "", ""])
sheet("客戶對照",
      ["整理版客戶(原字串)", "常見來源/單位", "出現次數", "累計銷售額",
       "建議客戶名", "建議分級", "★確認對應到(系統客戶名 / 新建)", "★合併到(若與他人同一人)", "備註"],
      data, [26, 18, 9, 11, 18, 10, 30, 24, 20], [7, 8])

# 通路對照
data = []
for (csrc, logi, plan), n in sorted(csrc_logi.items(), key=lambda x: -x[1]):
    data.append([csrc or "(空)", logi or "(空)", plan or "(空)", n,
                 guess_channel(csrc, logi, plan), "", ""])
sheet("通路對照",
      ["客戶來源/單位線索", "物流方式", "方案", "出現次數", "建議銷售管道", "★確認管道", "備註"],
      data, [24, 12, 12, 9, 14, 14, 20], [6])

# 批次對照
data = []
for season, n in sorted(season_cnt.items()):
    data.append([season, n, f"{season}-A", "", "", "", "整批可拆 A/B/C,依焙製日期"])
sheet("批次對照",
      ["產季", "訂單數", "建議批次代碼", "★產出數量", "★單位成本", "★製造日(YYYY-MM-DD)", "備註"],
      data, [8, 9, 14, 12, 12, 20, 28], [4, 5, 6])

# 價格參考(唯讀,不用填)
data = []
for (item, spec), prs in sorted(prod_price.items(), key=lambda x: x[0]):
    if not prs:
        continue
    data.append([item, spec or "(空)", len(prs), min(prs), max(prs),
                 Counter(prs).most_common(1)[0][0]])
sheet("價格參考",
      ["品項", "規格", "銷售筆數", "最低價", "最高價", "最常見價"],
      data, [12, 14, 10, 10, 10, 10], [])

# 說明
ws = wb2.create_sheet("說明", 0)
for line in [
    ["這份是「匯入對照表」——把整理版的舊資料對應到系統的主檔。"],
    ["黃底欄位(★開頭)請郡碩農創填寫;其餘為系統的建議值,可直接改。"],
    [""],
    ["商品對照", "整理版每個『品項＋規格』要對到系統哪個 SKU。系統沒有的規格,★確認SKU 填新代號並在備註寫清楚。"],
    ["客戶對照", "整理版每個客戶原字串要對到系統哪一位。同一人不同寫法,★合併到 填要保留的那個名字。查不到的填『新建』。"],
    ["通路對照", "每種『來源+物流』組合對到哪個銷售管道(官網 / 市集展售 / 批發 / 媒體導流 / LINE社群 / 自售…)。"],
    ["批次對照", "每個產季要建幾個批次、各批的產出量 / 單位成本 / 製造日。單位成本影響毛利計算。"],
    ["價格參考", "唯讀。列出每個品項規格過去實際成交價的分佈,供設定定價時參考。"],
    [""],
    ["填完後交回,再跑第二支腳本『匯入歷史.py』把資料寫進資料庫。"],
    ["整理版『資料檢核』有標記的 23 筆(金額不符 / 重複 / 重號等)不會匯入,另列清單。"],
] :
    ws.append(line)
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 90
for r in ws.iter_rows():
    r[0].font = Font(bold=True)
    if len(r) > 1:
        r[1].alignment = Alignment(wrap_text=True, vertical="top")

wb2.save(OUT)
print(f"已產出:{OUT}")
print(f"  商品組合 {len(prod)} 種 · 客戶 {len(cust_stat)} 個 · 來源/物流組合 {len(csrc_logi)} 種 · 產季 {len(season_cnt)} 個")
flagged = sum(1 for r in rows if r[C_CHECK])
print(f"  整理版有標記(不匯入)的列:{flagged}")
